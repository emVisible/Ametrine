#!/usr/bin/env python
"""格式基准第二步：**内容到底还剩多少进了知识库**。

上一轮只量了「检索准不准」，那默认了「文档本来就被完整索引了」。
这一条默认必须由独立手段验一遍：解析器自己说自己没错是不算数的，
所以参考正文由**另一套本地解析**（python-docx / python-pptx / openpyxl / pypdf /
email / csv / BeautifulSoup）取出，再和应用存回来的分块比。

单位（unit）按格式的自然结构定：段落 / CSV 行 / 幻灯片 / 表格行 / 邮件头字段 / PDF 页内句段。
判「这个单位在不在库里」用它的三个采样片段（头 25 / 中 30 / 尾 25 字符，空白归一后），
命中任一段即算在 —— 因为分块边界可能正好切在单位中间，只比整段会把这种情况误判成丢失。

用法：
    apps/backend/.venv/bin/python rag-bench/formats_fidelity.py --password '<基准账号口令>'
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from eval_retrieval import call, unwrap  # noqa: E402  同一套 HTTP/信封处理，不写第二份

ROOT = Path(__file__).resolve().parent
SKIP_TAGS = {"script", "style", "noscript", "svg", "nav", "header", "footer", "aside", "form"}


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip().lower()


def units_from_text(text: str, min_len: int = 60) -> list[str]:
    """按空行切段。段里可能带硬换行（RFC 正文每 72 列折行），
    所以必须和索引文本一样做空白归一 —— 不归一的话，采样片段里带着 `\n`，
    在归一过的 haystack 里永远找不到，「保真度 0.5」这种数字就全是我的读法造成的。
    """
    parts = [re.sub(r"[#*>`|\[\]-]+", " ", p).strip() for p in re.split(r"\n\s*\n", text)]
    return [norm(p) for p in parts if len(p) >= min_len]


def source_units(ext: str, blob: bytes) -> tuple[list[str], str]:
    """返回（该格式的自然单位列表, 用了什么抽取手段）。"""
    how = ""
    try:
        if ext in (".md", ".markdown", ".txt"):
            t = blob.decode("utf-8", "replace")
            how = "直接读文本"
            return units_from_text(t), how
        if ext == ".csv":
            rows = list(csv.reader(io.StringIO(blob.decode("utf-8", "replace"))))
            how = f"csv 模块（{len(rows)} 行）"
            # 不能用整行原文去比：CSVLoader 把每一行重写成 `列名: 值` 的形态，
            # 于是 `"USW00094728,2024-02-14,0.03"` 这种行串在索引文本里永远找不到 ——
            # 上一版就这样读出「.csv 保真度 0.0」，那完全是我的比法错了。
            # 表格的真问题只有一个：**每个单元格值有没有进库**（以及表头有没有留下）。
            out = [norm(str(c)) for r in rows for c in r if len(norm(str(c))) >= 6]
            return out, how
        if ext == ".html":
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(blob, "html.parser")
            for tag in soup(list(SKIP_TAGS)):
                tag.decompose()
            text = soup.get_text("\n")
            how = "bs4 去脚本/导航后的正文"
            return units_from_text(text), how
        if ext == ".eml":
            from email import message_from_bytes
            msg = message_from_bytes(blob)
            out = [norm(f"{h}: {msg[h]}") for h in ("From", "To", "Subject", "Date") if msg[h]]
            payload = msg.get_payload(decode=True) or b""
            body = payload.decode("utf-8", "replace")
            out += units_from_text(re.sub(r"<[^>]+>", " ", body), min_len=50)
            how = "email 模块（头字段 + 正文段）"
            return out, how
        if ext == ".docx":
            import docx
            d = docx.Document(io.BytesIO(blob))
            out = [norm(p.text) for p in d.paragraphs if len(p.text) >= 50]
            for table in d.tables:
                for row in table.rows:
                    cells = " ".join(c.text for c in row.cells)
                    if len(cells) >= 40:
                        out.append(norm(cells))
            how = f"python-docx（{len(d.paragraphs)} 段 / {len(d.tables)} 表）"
            return out, how
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(io.BytesIO(blob))
            out, n_slides = [], 0
            for s in prs.slides:
                n_slides += 1
                texts = []
                for sh in s.shapes:
                    if sh.has_text_frame:
                        texts.append(sh.text_frame.text)
                joined = norm(" ".join(texts))
                if len(joined) >= 40:
                    out.append(joined)
            how = f"python-pptx（{n_slides} 页）"
            return out, how
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
            out, nrows = [], 0
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = " ".join(str(c) for c in row if c is not None)
                    nrows += 1
                    if len(cells) >= 40:
                        out.append(norm(cells))
            how = f"openpyxl（{len(wb.worksheets)} 表 / {nrows} 行）"
            return out, how
        if ext == ".odt":
            with zipfile.ZipFile(io.BytesIO(blob)) as z:
                xml = z.read("content.xml").decode("utf-8", "replace")
            paras = re.findall(r"<text:p[^>]*>(.*?)</text:p>", xml, re.S)
            out = [norm(re.sub(r"<[^>]+>", "", p)) for p in paras]
            how = "content.xml 的 text:p"
            return [p for p in out if len(p) >= 50], how
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(blob))
            out = []
            for page in reader.pages:
                try:
                    t = page.extract_text() or ""
                except Exception:  # noqa: BLE001
                    t = ""
                out += [norm(s) for s in re.split(r"(?<=[.。])\s+", t) if len(s) >= 60]
            how = f"pypdf（{len(reader.pages)} 页文本层）"
            return out, how
    except Exception as e:  # noqa: BLE001
        return [], f"参考抽取失败：{type(e).__name__}: {e}"
    return [], f"未支持参考抽取的格式 {ext}"


def present(unit: str, haystack: str) -> bool:
    """单位是否进了索引文本。

    短单位（CSV 的一个单元格、日期、站号）必须**整段比**：
    原来只接受长度 ≥12 的采样片段，于是 11 个字符的站号永远算「没找到」，
    覆盖率被我的判据压成 0 —— 又是一个「指标把故障读成差成绩」的样子。
    """
    if not unit:
        return True
    if len(unit) < 12:
        return unit in haystack
    probes = [unit[:25], unit[len(unit) // 2:len(unit) // 2 + 30], unit[-25:]]
    return any(len(p) >= 12 and p in haystack for p in probes)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--password", required=True)
    ap.add_argument("--manifest", default="formats/MANIFEST.json")
    ap.add_argument("--map", default="ingest_formats_map.json", dest="map_name")
    ap.add_argument("--out", default="formats-fidelity.json")
    ap.add_argument("--report", default="formats-fidelity.md")
    args = ap.parse_args()

    manifest = json.loads((ROOT / args.manifest).read_text(encoding="utf-8"))
    mapping = json.loads((ROOT / args.map_name).read_text(encoding="utf-8"))
    st, res = call("/auth", form={"username": "ragbench", "password": args.password}, method="POST")
    if st != 200:
        print(f"登录失败 {st}: {str(res)[:160]}")
        return 1
    tok = res["access_token"]

    rows = []
    for doc in manifest["documents"]:
        ent = mapping.get(doc["id"])
        if not ent:
            rows.append({**{k: doc[k] for k in ("id", "format", "domain", "title")},
                         "ingested": False, "note": "没有入库映射"})
            continue
        # method="GET" 不是风格问题：`call()` 的默认方法是 POST，
        # POST 到这条只读的接口会拿到一个错误信封，`unwrap` 之后变成空列表 ——
        # 于是「0 块、保真度 0.0」全是我读法错的产物。
        st, res = call(f"/relation/document/chunk?doc_id={ent['doc_id']}",
                       tok=tok, method="GET")
        chunks = unwrap(res) or []
        if not isinstance(chunks, list):
            rows.append({"id": doc["id"], "format": doc["format"], "ingested": True,
                         "note": f"分块读取异常 {st}: {str(chunks)[:80]}"})
            continue
        # 字段名是 `content`，不是 `text`（实测：按 text 取会每篇都读到空串，
        # 于是「保真度 0」这种结论会完全来自我自己的读法）。两个名字都认，
        # 以后服务端改名时这条测量不会静默变成 0。
        def chunk_text(c: dict) -> str:
            for key in ("content", "text", "page_content"):
                if isinstance(c.get(key), str):
                    return c[key]
            return ""

        indexed = norm(" ".join(chunk_text(c) for c in chunks))
        blob = (ROOT / doc["path"]).read_bytes()
        units, how = source_units(doc["format"], blob)
        hit = sum(1 for u in units if present(u, indexed))
        src_chars = sum(len(u) for u in units)
        rows.append({
            "id": doc["id"], "format": doc["format"], "domain": doc["domain"],
            "title": doc["title"], "container": doc.get("container", ""),
            "bytes": doc["bytes"], "ingested": True, "doc_id": ent["doc_id"],
            "chunks": len(chunks),
            "enabled_chunks": sum(1 for c in chunks if c.get("enabled") is not False),
            "indexed_chars": len(indexed),
            "source_units": len(units), "units_found": hit,
            "unit_coverage": (round(hit / len(units), 3) if units else None),
            "extractor": how,
            "source_sample": units[0][:90] if units else "",
            "indexed_sample": indexed[:90],
        })
        print(f"  {doc['format']:<6} {doc['id'][:28]:<30} 块{len(chunks):<4} "
              f"单位{hit}/{len(units)} 索引字数{len(indexed):,}")

    (ROOT / args.out).write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")

    by_fmt: dict[str, list[dict]] = {}
    for r in rows:
        by_fmt.setdefault(r["format"], []).append(r)
    lines = ["# 多格式入库保真度", "",
             "参考正文由本地另一套解析器取出，再和应用存回来的分块比对。",
             "`unit_coverage` = 该格式的自然单位（段落/行/页/表行）里，有多少能在索引文本中找到。",
             "", "| 格式 | 份数 | 入库成功 | 平均块数 | 平均单位覆盖率 | 平均索引字数 | 抽取手段 |",
             "| --- | --- | --- | --- | --- | --- | --- |"]
    for fmt in sorted(by_fmt):
        rs = by_fmt[fmt]
        ok = [r for r in rs if r.get("chunks") is not None and r.get("source_units") is not None]
        cov = [r["unit_coverage"] for r in ok if r.get("unit_coverage") is not None]
        nch = [r["chunks"] for r in ok]
        chars = [r["indexed_chars"] for r in ok]
        ingested = sum(1 for r in rs if r.get("ingested") and r.get("chunks"))
        lines.append(
            f"| `{fmt}` | {len(rs)} | {ingested} | "
            f"{(round(sum(nch) / len(nch), 1) if nch else '-')} | "
            f"{(round(sum(cov) / len(cov), 3) if cov else '-')} | "
            f"{(round(sum(chars) / len(chars)) if chars else '-')} | "
            f"{(ok[0]['extractor'] if ok else '—')} |")
    fails = [r for r in rows if not r.get("ingested") or not r.get("chunks")]
    lines += ["", f"**没拿到分块的文档**：{len(fails)} 份", ""]
    for r in fails[:20]:
        lines.append(f"- `{r.get('format')}` {r['id'][:34]}：{str(r.get('note') or '入库成功但没有分块')[:120]}")
    low = [r for r in rows if r.get("unit_coverage") is not None and r["unit_coverage"] < 0.6]
    if low:
        lines += ["", "**覆盖率低于 0.6 的（内容在里面却查不到的那一类）**", "",
                  "| 格式 | 文档 | 覆盖率 | 块数 | 索引字数 | 单位数 |", "| --- | --- | --- | --- | --- | --- |"]
        for r in sorted(low, key=lambda x: x["unit_coverage"]):
            lines.append(f"| `{r['format']}` | {r['id'][:30]} | {r['unit_coverage']} | "
                         f"{r['chunks']} | {r['indexed_chars']:,} | {r['source_units']} |")
    (ROOT / args.report).write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n" + "\n".join(lines))
    return 0


if __name__ == "__main__":
    sys.exit(main())
